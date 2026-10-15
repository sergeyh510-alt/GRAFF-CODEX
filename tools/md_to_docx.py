#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Markdown -> .docx converter for the GRAFF articles.

Turns a Markdown file into a clean Word document with python-docx:

* ``#`` / ``##`` / ``###`` map to the Title / Heading 1 / Heading 2 styles;
* paragraphs support inline ``**bold**``, ``*italic*`` and ``code``;
* ``-`` / ``*`` / ``+`` bullets and ``1.`` items become list paragraphs;
* pipe tables (first row = header) become a real Word table;
* ``![alt](path)`` inserts the image scaled to the text width (never wider
  than ~16 cm, aspect ratio preserved); the path is resolved relative to the
  Markdown file;
* ``mermaid`` fenced blocks are skipped (the diagrams ship as PNG images
  already referenced in the article); any other fenced block is emitted in a
  monospaced font.

Usage::

    python tools/md_to_docx.py                       # both shipped articles
    python tools/md_to_docx.py IN.md OUT.docx        # a single pair
"""

from __future__ import annotations

import os
import re
import struct
import sys

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

# --- layout constants ------------------------------------------------------ #
BODY_FONT = "Calibri"
BODY_SIZE = Pt(11)
MONO_FONT = "Consolas"
CODE_SIZE = Pt(10)

MAX_IMAGE_WIDTH_CM = 16.0     # no image is wider than the text column
MAX_IMAGE_HEIGHT_CM = 18.0    # ... and none runs off the page either
TABLE_STYLE = "Table Grid"

# PNGs saved without a pHYs chunk carry no DPI; Word would then render them at
# a huge default size, so assume the usual screen DPI and scale from pixels.
FALLBACK_DPI = 96.0

REPO_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

DEFAULT_JOBS = [
    (os.path.join(REPO_ROOT, "docs", "linkedin-article-en.md"),
     os.path.join(REPO_ROOT, "docs", "GRAFF_LinkedIn_EN.docx")),
    (os.path.join(REPO_ROOT, "docs", "linkedin-article-ru.md"),
     os.path.join(REPO_ROOT, "docs", "GRAFF_LinkedIn_RU.docx")),
]

# --- inline formatting ----------------------------------------------------- #
_TOKEN_RE = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*|\*[^*]+\*)")


def add_runs(paragraph, text):
    """Append *text* to *paragraph*, honouring **bold**, *italic* and `code`."""
    for token in _TOKEN_RE.split(text):
        if not token:
            continue
        if token.startswith("`") and token.endswith("`") and len(token) > 2:
            run = paragraph.add_run(token[1:-1])
            run.font.name = MONO_FONT
            run.font.size = CODE_SIZE
        elif token.startswith("**") and token.endswith("**") and len(token) > 4:
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("*") and token.endswith("*") and len(token) > 2:
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        else:
            paragraph.add_run(token)


# --- image sizing ---------------------------------------------------------- #
def png_meta(path):
    """Return (width_px, height_px, dpi_or_None) for a PNG, else None."""
    try:
        with open(path, "rb") as fh:
            data = fh.read()
    except OSError:
        return None
    if len(data) < 24 or data[:8] != b"\x89PNG\r\n\x1a\n":
        return None
    width, height = struct.unpack(">II", data[16:24])
    dpi = None
    idx = data.find(b"pHYs")
    if idx != -1 and idx + 13 <= len(data):
        px = int.from_bytes(data[idx + 4:idx + 8], "big")
        unit = data[idx + 12]
        if unit == 1 and px > 0:
            dpi = px * 0.0254
    return width, height, dpi


def scaled_image_size(path):
    """(width_cm, height_cm) for the image, capped and ratio-preserving."""
    meta = png_meta(path)
    if not meta:
        return None
    width_px, height_px, dpi = meta
    if not width_px or not height_px:
        return None
    dpi = dpi or FALLBACK_DPI
    width_cm = width_px / dpi * 2.54
    height_cm = height_px / dpi * 2.54
    scale = 1.0
    if width_cm > MAX_IMAGE_WIDTH_CM:
        scale = MAX_IMAGE_WIDTH_CM / width_cm
    if height_cm * scale > MAX_IMAGE_HEIGHT_CM:
        scale = MAX_IMAGE_HEIGHT_CM / height_cm
    return width_cm * scale, height_cm * scale


# --- block builders -------------------------------------------------------- #
def configure_document(doc):
    """Page margins, base Calibri 11 font and comfortable paragraph spacing."""
    section = doc.sections[0]
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = BODY_SIZE
    try:  # make the font stick for East-Asian text too
        normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    except Exception:
        pass
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.15


def add_heading(doc, level, text):
    if level <= 1:
        style = "Title"
    elif level == 2:
        style = "Heading 1"
    elif level == 3:
        style = "Heading 2"
    elif level == 4:
        style = "Heading 3"
    else:
        style = "Heading 4"
    paragraph = doc.add_paragraph(style=style)
    add_runs(paragraph, text)
    return paragraph


def add_list_item(doc, text, numbered):
    style = "List Number" if numbered else "List Bullet"
    paragraph = doc.add_paragraph(style=style)
    add_runs(paragraph, text)
    return paragraph


def add_code_block(doc, lines):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(4)
    paragraph.paragraph_format.space_after = Pt(10)
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run("\n".join(lines))
    run.font.name = MONO_FONT
    run.font.size = CODE_SIZE
    return paragraph


def add_image(doc, base_dir, rel_path):
    path = rel_path if os.path.isabs(rel_path) else os.path.normpath(os.path.join(base_dir, rel_path))
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(10)
    if not os.path.exists(path):
        paragraph.add_run("[missing image: %s]" % rel_path)
        return paragraph
    size = scaled_image_size(path)
    run = paragraph.add_run()
    if size:
        run.add_picture(path, width=Cm(size[0]), height=Cm(size[1]))
    else:
        run.add_picture(path, width=Cm(MAX_IMAGE_WIDTH_CM))
    return paragraph


def parse_table_row(line):
    text = line.strip()
    if text.startswith("|"):
        text = text[1:]
    if text.endswith("|"):
        text = text[:-1]
    return [cell.strip() for cell in text.split("|")]


def is_separator_row(cells):
    return bool(cells) and all(re.fullmatch(r":?-{2,}:?", c or "") for c in cells)


def add_table(doc, rows):
    header, body = rows[0], rows[1:]
    table = doc.add_table(rows=len(rows), cols=len(header))
    try:
        table.style = TABLE_STYLE
    except KeyError:
        pass
    for col, text in enumerate(header):
        paragraph = table.cell(0, col).paragraphs[0]
        add_runs(paragraph, text)
        for run in paragraph.runs:
            run.bold = True
    for row_idx, row in enumerate(body, start=1):
        for col in range(len(header)):
            paragraph = table.cell(row_idx, col).paragraphs[0]
            add_runs(paragraph, row[col] if col < len(row) else "")
    doc.add_paragraph()  # a little air after the table
    return table


# --- driver ---------------------------------------------------------------- #
FENCE_RE = re.compile(r"^(`{3,}|~{3,})(.*)$")
HEADING_RE = re.compile(r"^(#{1,6})\s+(.*)$")
IMAGE_RE = re.compile(r"^!\[(.*?)\]\(\s*(\S+?)\s*\)\s*$")
BULLET_RE = re.compile(r"^[-*+]\s+(.*)$")
NUMBER_RE = re.compile(r"^\d+[.)]\s+(.*)$")


def convert(md_path, docx_path):
    with open(md_path, encoding="utf-8") as fh:
        lines = fh.read().split("\n")
    base_dir = os.path.dirname(os.path.abspath(md_path))

    doc = Document()
    configure_document(doc)

    i = 0
    total = len(lines)
    while i < total:
        line = lines[i]
        stripped = line.strip()

        fence = FENCE_RE.match(line)
        if fence:
            info = fence.group(2).strip().lower()
            code_lines = []
            i += 1
            while i < total and not FENCE_RE.match(lines[i]):
                code_lines.append(lines[i])
                i += 1
            i += 1  # consume the closing fence
            if info.startswith("mermaid"):
                continue  # the diagram is shipped as a PNG already
            add_code_block(doc, code_lines)
            continue

        if not stripped:
            i += 1
            continue

        heading = HEADING_RE.match(stripped)
        if heading:
            add_heading(doc, len(heading.group(1)), heading.group(2).strip())
            i += 1
            continue

        image = IMAGE_RE.match(stripped)
        if image:
            add_image(doc, base_dir, image.group(2))
            i += 1
            continue

        if stripped.startswith("|"):
            rows = []
            while i < total and lines[i].strip().startswith("|"):
                rows.append(parse_table_row(lines[i]))
                i += 1
            rows = [r for r in rows if not is_separator_row(r)]
            add_table(doc, rows)
            continue

        bullet = BULLET_RE.match(stripped)
        if bullet:
            add_list_item(doc, bullet.group(1), numbered=False)
            i += 1
            continue

        number = NUMBER_RE.match(stripped)
        if number:
            add_list_item(doc, number.group(1), numbered=True)
            i += 1
            continue

        paragraph = doc.add_paragraph()
        add_runs(paragraph, stripped)
        i += 1

    os.makedirs(os.path.dirname(os.path.abspath(docx_path)), exist_ok=True)
    doc.save(docx_path)
    return docx_path


def main(argv):
    if not argv:
        jobs = DEFAULT_JOBS
    elif len(argv) == 2:
        jobs = [(argv[0], argv[1])]
    else:
        print(__doc__)
        return 2
    for src, dst in jobs:
        convert(src, dst)
        print("wrote %s" % dst)
    return 0


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
